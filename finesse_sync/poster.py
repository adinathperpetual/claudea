"""Upload extracted transactions to Finesse and post them from Equity Staging.

Step 1 (automatic, ``start``):
    File Upload > Equity Uploads > Transactions Upload > "Excel Upload" > attach the
    filled template > Upload; then Equity > Equity Staging: search each trading
    account in the file and find exactly the rows this upload created (same trade
    date, buy/sell, quantity and rate). Rows already identical in staging *before*
    the upload are not counted (and are reported as a possible duplicate upload).
    The job then waits for a person to confirm.

Step 2 (``confirm``, after the person clicks "Confirm & Post"):
    For each account: search, open "Mapped", tick only the planned rows, Post,
    accept the confirmation, and check that they left staging. Unmapped rows are
    never posted; they are reported so they can be fixed in Finesse.
"""
from __future__ import annotations

import io
import json
import logging
import re
import threading
import time
from dataclasses import asdict, dataclass, field
from datetime import date, datetime

from . import db
from .config import get_settings, redact
from .finesse import BrowserSession, FinesseClient, FinesseError
from .records import LayoutChangedError

log = logging.getLogger("finesse_sync.poster")

_lock = threading.Lock()          # one upload/post job at a time
MSG_SEL = ("simple-snack-bar, .mat-mdc-snack-bar-container, mat-snack-bar-container, .toast, .toast-message, "
           "[role=alert], [role=status], mat-dialog-container, .swal2-popup, .alert, .notification")


class PostError(RuntimeError):
    pass


# ---------------------------------------------------------------- rows
@dataclass
class TxnRow:
    account: str
    client: str
    trade_date: str       # ISO yyyy-mm-dd
    side: str             # PURCHASE / SELL
    qty: float
    rate: float
    isin: str = ""

    def key(self) -> tuple:
        return (self.account.upper(), self.trade_date, self.side, round(self.qty, 4), round(self.rate, 2))

    def label(self) -> str:
        return f"{self.trade_date} {self.side} {self.qty:g} @ {self.rate:.2f} {self.isin}".strip()


def _num(v) -> float | None:
    try:
        s = str(v).replace(",", "").replace("₹", "").strip()
        return float(s) if s not in ("", "-", "None") else None
    except ValueError:
        return None


def _iso(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%Y/%m/%d", "%d-%m-%y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    m = re.match(r"^(\d+(?:\.\d+)?)$", s)          # Excel serial date
    if m:
        return date.fromordinal(date(1899, 12, 30).toordinal() + int(float(m.group(1)))).isoformat()
    return s


def _side(v) -> str:
    s = str(v or "").strip().upper()
    if s in ("BUY", "B", "PURCHASE", "P"):
        return "PURCHASE"
    if s in ("SELL", "S", "SALE"):
        return "SELL"
    return s


def rows_from_template(data: bytes) -> list[TxnRow]:
    """Read the filled template (Trading Account | Client Name | Trade Date | … | ISIN No |
    Trade Type | Transaction Type | Quantity | Market Price Per Share | …)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise PostError("The file is empty.")
    head = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            for i, h in enumerate(head):
                if h == n:
                    return i
        raise PostError(f"Column {names[0]!r} not found in the upload file (columns: {head}).")

    a, c, d = col("trading account"), col("client name"), col("trade date")
    t, q, p = col("transaction type"), col("quantity"), col("market price per share")
    i = col("isin no", "isin")
    out = []
    for r in rows[1:]:
        if not r or all(v in (None, "") for v in r):
            continue
        qty, rate = _num(r[q]), _num(r[p])
        if not r[a] or qty is None or rate is None:
            raise PostError(f"Row for {r[c] or '?'} is missing the trading account, quantity or price — "
                            "fix it in the extractor before uploading.")
        out.append(TxnRow(str(r[a]).strip().upper(), str(r[c] or "").strip(), _iso(r[d]), _side(r[t]),
                          qty, rate, str(r[i] or "").strip()))
    if not out:
        raise PostError("No transactions in the upload file.")
    return out


# ---------------------------------------------------------------- jobs (persisted)
SCHEMA = """CREATE TABLE IF NOT EXISTS post_jobs (
    id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
    created_by TEXT, file_name TEXT, status TEXT NOT NULL, step TEXT, message TEXT, data TEXT)"""


_ready = False


def _init() -> None:
    """Create the table; on the first call in this process, close jobs a restart interrupted."""
    global _ready
    with db.connect() as c:
        c.execute(SCHEMA)
        if not _ready:
            c.execute("UPDATE post_jobs SET status='failed', message='Interrupted (service restarted)' "
                      "WHERE status IN ('uploading', 'posting')")
    _ready = True


def _save(job_id: int, **kw) -> None:
    sets, args = ["updated_at=?"], [db.now_iso()]
    for k, v in kw.items():
        sets.append(f"{k}=?")
        args.append(json.dumps(v) if k == "data" else v)
    with db.connect() as c:
        c.execute(f"UPDATE post_jobs SET {', '.join(sets)} WHERE id=?", (*args, job_id))


def get_job(job_id: int) -> dict | None:
    _init()
    with db.connect() as c:
        r = c.execute("SELECT * FROM post_jobs WHERE id=?", (job_id,)).fetchone()
    if not r:
        return None
    d = dict(r)
    d["data"] = json.loads(d["data"] or "{}")
    return d


def list_jobs(limit: int = 20) -> list[dict]:
    _init()
    with db.connect() as c:
        rows = c.execute("SELECT id FROM post_jobs ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    return [get_job(r["id"]) for r in rows]


# ---------------------------------------------------------------- auto-learn
# What worked on the last run (the search box, the confirm button …) is remembered in
# data/finesse_learned.json and tried first next time, so the tool goes straight there
# instead of scanning the page. Anything that no longer works is simply re-discovered.
def _learned_path():
    return get_settings().data_dir / "finesse_learned.json"


def learned(key: str, default=None):
    try:
        return json.loads(_learned_path().read_text(encoding="utf-8")).get(key, default)
    except (OSError, ValueError):
        return default


def learn(key: str, value) -> None:
    try:
        p = _learned_path()
        d = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}
        if d.get(key) != value:
            d[key] = value
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_text(json.dumps(d, indent=1), encoding="utf-8")
    except (OSError, ValueError):
        pass


# ---------------------------------------------------------------- jobs
REVIEW_MINUTES = 20        # how long the Finesse browser stays open while you review
_waiters: dict[int, dict] = {}


def start(data: bytes, file_name: str, user: str = "") -> int:
    """Upload the file and find the new rows in Equity Staging; then wait (browser kept
    open) for confirm()/cancel(). Nothing is posted without confirm()."""
    rows = rows_from_template(data)        # validate before touching Finesse
    _init()
    if not _lock.acquire(blocking=False):
        raise PostError("Another upload to Finesse is running — wait for it to finish.")
    ts = db.now_iso()
    with db.connect() as c:
        job_id = int(c.execute(
            "INSERT INTO post_jobs (created_at, updated_at, created_by, file_name, status, step, data) "
            "VALUES (?,?,?,?,?,?,?)", (ts, ts, user, file_name, "uploading", "Starting", json.dumps(
                {"rows": [asdict(r) for r in rows]}))).lastrowid)
    waiter = {"event": threading.Event(), "action": None, "selected": None}
    _waiters[job_id] = waiter

    def run():
        try:
            def work(b: BrowserSession):
                _phase_upload(job_id, b, data, file_name, rows)
                job = get_job(job_id)
                if job["status"] != "awaiting_confirmation":
                    return
                # keep this logged-in browser on Equity Staging while the person reviews
                if waiter["event"].wait(REVIEW_MINUTES * 60) and waiter["action"] == "confirm":
                    _phase_post(job_id, b, get_job(job_id)["data"], waiter["selected"])
            _with_browser(work)
        except Exception as e:  # noqa: BLE001
            log.exception("Upload to Finesse failed")
            _save(job_id, status="failed", message=redact(str(e))[:1000])
        finally:
            _waiters.pop(job_id, None)
            _lock.release()

    threading.Thread(target=run, name=f"finesse-upload-{job_id}", daemon=True).start()
    return job_id


def _check_selection(job: dict, selected: list[str] | None) -> list[str]:
    cands = {c["id"]: c for a in job["data"]["accounts"].values() for c in a["candidates"]}
    if selected is None:
        selected = [i for i, c in cands.items() if c["kind"] == "new"]
    bad = [i for i in selected if i not in cands or cands[i]["kind"] not in ("new", "duplicate")]
    if bad:
        raise PostError("These rows cannot be posted: " + ", ".join(bad))
    if not selected:
        raise PostError("No rows selected.")
    return selected


def confirm(job_id: int, selected: list[str] | None = None) -> None:
    """Post the selected rows (ids from the job's candidates; default: all new rows)."""
    job = get_job(job_id)
    if not job:
        raise PostError("Job not found.")
    if job["status"] != "awaiting_confirmation":
        raise PostError(f"This job is '{job['status']}', not waiting for confirmation.")
    selected = _check_selection(job, selected)
    _save(job_id, status="posting", step="Starting")
    w = _waiters.get(job_id)
    if w:                                   # the browser is still open on Equity Staging
        w["selected"], w["action"] = selected, "confirm"
        w["event"].set()
        return
    if not _lock.acquire(blocking=False):
        raise PostError("Another upload to Finesse is running — wait for it to finish.")

    def run():
        try:
            _with_browser(lambda b: _phase_post(job_id, b, job["data"], selected))
        except Exception as e:  # noqa: BLE001
            log.exception("Posting in Finesse failed")
            _save(job_id, status="failed", message=redact(str(e))[:1000])
        finally:
            _lock.release()

    threading.Thread(target=run, name=f"finesse-post-{job_id}", daemon=True).start()


def cancel(job_id: int) -> None:
    job = get_job(job_id)
    if not job:
        raise PostError("Job not found.")
    if job["status"] != "awaiting_confirmation":
        raise PostError(f"This job is '{job['status']}' and cannot be cancelled now.")
    _save(job_id, status="cancelled",
          message="Not posted. The uploaded rows are still in Equity Staging — delete them there if not needed.")
    w = _waiters.get(job_id)
    if w:
        w["action"] = "cancel"
        w["event"].set()


# ---------------------------------------------------------------- browser steps
def _with_browser(work):
    fc = FinesseClient()
    try:
        return fc._in_browser(work)
    finally:
        fc.close()


def _messages(b: BrowserSession) -> list[str]:
    out = []
    try:
        loc = b.page.locator(MSG_SEL)
        for i in range(min(loc.count(), 10)):
            el = loc.nth(i)
            if el.is_visible():
                t = re.sub(r"\s+", " ", el.inner_text() or "").strip()
                if t and t not in out:
                    out.append(t[:300])
    except Exception:  # noqa: BLE001
        pass
    return out


def _wait_message(b: BrowserSession, seconds: float) -> str:
    deadline = time.time() + seconds
    while time.time() < deadline:
        msgs = [m for m in _messages(b) if re.search(r"success|upload|error|fail|invalid|posted|saved|done|complete", m, re.I)]
        if msgs:
            return " | ".join(msgs)
        b.page.wait_for_timeout(250)
    return ""


def _click_button(b: BrowserSession, rx: str, exclude: str = r"template|download", learn_key: str = "") -> bool:
    if learn_key:                                   # the button that worked last time
        txt = learned(learn_key)
        if txt:
            loc = b.page.get_by_role("button", name=re.compile(rf"^\s*{re.escape(txt)}\s*$", re.I))
            for i in range(min(loc.count(), 5)):
                el = loc.nth(loc.count() - 1 - i)
                if el.is_visible() and not el.is_disabled():
                    el.click()
                    return True
    btns = b.page.locator("button, [role=button], input[type=submit], input[type=button]")
    found, found_txt = None, ""
    for i in range(min(btns.count(), 80)):
        el = btns.nth(i)
        try:
            if not el.is_visible() or el.is_disabled():
                continue
            t = re.sub(r"\s+", " ", (el.inner_text() or el.get_attribute("value") or "")).strip()
            if re.search(rx, t, re.I) and not re.search(exclude, t, re.I):
                found, found_txt = el, t      # last match: the action button sits below the headings
        except Exception:  # noqa: BLE001
            continue
    if found is None:
        return False
    found.click()
    if learn_key and found_txt:
        learn(learn_key, found_txt)
    return True


def _close_dialog(b: BrowserSession, labels: tuple, wait_s: float = 2, learn_key: str = "") -> bool:
    """Click the first matching button inside a visible dialog (never elsewhere on the page)."""
    first = learned(learn_key) if learn_key else None
    labels = ((first,) if first else ()) + tuple(l for l in labels if l != first)
    deadline = time.time() + wait_s
    while time.time() < deadline:
        dlg = b.page.locator("mat-dialog-container, .swal2-popup, [role=dialog], [role=alertdialog], .modal.show")
        for i in range(min(dlg.count(), 3)):
            d = dlg.nth(i)
            if not d.is_visible():
                continue
            for t in labels:
                btn = d.get_by_role("button", name=re.compile(rf"^\s*{re.escape(t)}\s*$", re.I))
                if btn.count() and btn.first.is_visible():
                    btn.first.click()
                    if learn_key:
                        learn(learn_key, t)
                    b.page.wait_for_timeout(300)
                    return True
        b.page.wait_for_timeout(150)
    return False


def _upload_file(b: BrowserSession, data: bytes, file_name: str) -> str:
    s = b.s
    b._goto(s.base_url + s.finesse_txn_upload_route, "Open Equity Transaction Upload")
    if not b._click_text("Excel Upload", wait_s=20):
        raise b._report_error("'Excel Upload' option not found on the Equity Transaction Upload page.")
    inputs = b.page.locator("input[type=file]")
    if inputs.count() == 0:
        raise b._report_error("No file box found on the Equity Transaction Upload page.")
    inputs.first.set_input_files(files=[{"name": file_name, "buffer": data,
                                         "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}])
    b.page.wait_for_timeout(300)
    if not _click_button(b, r"(^|\s)upload\s*$", learn_key="upload_button"):
        raise b._report_error("'Upload' button not found on the Equity Transaction Upload page.")
    msg = _wait_message(b, 120)
    if re.search(r"error|fail|invalid", msg, re.I) and not re.search(r"success", msg, re.I):
        raise b._report_error(f"Finesse rejected the upload: {msg}")
    _close_dialog(b, ("OK", "Ok", "Close", "Done"), wait_s=1, learn_key="upload_dialog")
    return msg or "Upload sent (Finesse showed no message)."


def _table_sig(b: BrowserSession) -> str:
    try:
        return b.page.evaluate("""() => { const t = [...document.querySelectorAll('table')]
              .find(t => /trade\\s*date/i.test(t.innerText)); return t ? t.innerText.length + ':' + t.innerText.slice(0, 400) : ''; }""")
    except Exception:  # noqa: BLE001
        return ""


def _wait_table(b: BrowserSession, before: str, ok=lambda rows: True, max_s: float = 20) -> list[dict]:
    """Return as soon as the staging grid has changed (or holds what ``ok`` wants) and
    stays the same for a moment — no fixed pauses."""
    t0 = time.time()
    deadline, last, stable_since = t0 + max_s, None, t0
    while time.time() < deadline:
        sig = _table_sig(b)
        if sig != last:
            last, stable_since = sig, time.time()
        elif time.time() - stable_since >= 0.4:
            # changed grid: accept once settled; unchanged grid (same result again): give the
            # page 1.5s to refresh before trusting it
            if sig != before or time.time() - t0 >= 1.5:
                rows = _staging_rows(b)
                if ok(rows):
                    return rows
        b.page.wait_for_timeout(120)
    return _staging_rows(b)


def _account_box(b: BrowserSession):
    css = learned("staging_box")                  # the box that worked last time
    if css:
        loc = b.page.locator(css)
        try:
            loc.first.wait_for(state="visible", timeout=15000)
            return loc.first
        except Exception:  # noqa: BLE001
            pass
    deadline = time.time() + 30                # the page draws itself a moment after it opens
    while time.time() < deadline:
        for f in b._visible_inputs():
            if re.search(r"trading\s*account", f["words"], re.I):
                a = f["attrs"]
                sel = (f"#{a['id']}" if a["id"] and re.fullmatch(r"[A-Za-z][\w-]*", a["id"]) else
                       f"input[formcontrolname='{a['formcontrolname']}']" if a["formcontrolname"] else
                       f"input[name='{a['name']}']" if a["name"] else "")
                if sel:
                    learn("staging_box", sel)
                return f["el"]
        b.page.wait_for_timeout(250)
    raise b._report_error("'Trading Account Number' box not found on Equity Staging.")


def _staging_search(b: BrowserSession, account: str) -> list[dict]:
    """Search one trading account; returns the rows shown (by the time the grid settled)."""
    box = _account_box(b)
    before = _table_sig(b)
    box.fill(account)
    if not _click_button(b, r"^\s*search\s*$", exclude=r"^$", learn_key="search_button"):
        box.press("Enter")
    return _wait_table(b, before, ok=lambda rows: all(not r["account"] or r["account"] == account.upper() for r in rows))


def _tab_count(b: BrowserSession, name: str) -> int | None:
    loc = b.page.get_by_text(re.compile(rf"^\s*{name}\s*\(\s*\d+\s*\)\s*$", re.I))
    for i in range(min(loc.count(), 5)):
        if loc.nth(i).is_visible():
            return int(re.search(r"\((\d+)\)", loc.nth(i).inner_text()).group(1))
    return None


def _tab(b: BrowserSession, name: str) -> list[dict]:
    """Open a staging tab like 'Mapped(3)' and return its rows once the grid has switched."""
    loc = b.page.get_by_text(re.compile(rf"^\s*{name}\s*\(\s*\d+\s*\)\s*$", re.I))
    for i in range(min(loc.count(), 5)):
        el = loc.nth(i)
        if el.is_visible():
            n = int(re.search(r"\((\d+)\)", el.inner_text()).group(1))
            before = _table_sig(b)
            el.click()
            return _wait_table(b, before, ok=lambda rows: len(rows) == n, max_s=8)
    return _staging_rows(b)


def _views(b: BrowserSession, rows_now: list[dict]) -> tuple[list[dict], list[dict]]:
    """(all rows, mapped rows) for the account on screen. When Finesse shows 0 unmapped,
    both are what is already on screen — no tab clicks needed."""
    unmapped = (_tab_count(b, "Unmapped Scrips") or 0) + (_tab_count(b, "Unmapped Clients") or 0)
    if unmapped == 0:
        return rows_now, rows_now
    return _tab(b, "All"), _tab(b, "Mapped")


def _staging_rows(b: BrowserSession) -> list[dict]:
    """Rows of the staging table as dicts (by header text) + their position."""
    data = b.page.evaluate("""() => {
      const t = [...document.querySelectorAll('table')].find(t => /trade\\s*date/i.test(t.innerText) && /quantity/i.test(t.innerText));
      if (!t) return null;
      const head = [...t.querySelectorAll('thead th, thead td')].map(c => c.innerText.replace(/\\s+/g, ' ').trim().toLowerCase());
      const rows = [...t.querySelectorAll('tbody tr')].filter(r => r.querySelector('td'));
      document.querySelectorAll('[data-cne-pos]').forEach(r => r.removeAttribute('data-cne-pos'));
      rows.forEach((r, i) => r.setAttribute('data-cne-pos', String(i)));
      return { head, rows: rows.map(r => [...r.querySelectorAll('td')].map(c => c.innerText.replace(/\\s+/g, ' ').trim())) };
    }""")
    if not data:
        return []
    head, out = data["head"], []

    def idx(rx):
        return next((i for i, h in enumerate(head) if re.search(rx, h)), None)

    ia, idt, it, iq = idx(r"trading\s*account"), idx(r"trade\s*date"), idx(r"transaction\s*type"), idx(r"quantity")
    imr, inr, isc = idx(r"market\s*rate"), idx(r"net\s*rate"), idx(r"scrip")
    for pos, r in enumerate(data["rows"]):
        if len(r) < len(head) - 1 or idt is None or iq is None:
            continue
        get = lambda i: r[i] if i is not None and i < len(r) else ""  # noqa: E731
        out.append({"pos": pos, "account": get(ia).upper(), "date": _iso(get(idt)), "side": _side(get(it)),
                    "qty": _num(get(iq)), "market": _num(get(imr)), "net": _num(get(inr)), "scrip": get(isc)})
    return out


def _matches(row: dict, t: TxnRow) -> bool:
    if row["date"] != t.trade_date or row["side"] != t.side or row["qty"] is None or abs(row["qty"] - t.qty) > 1e-6:
        return False
    if row["account"] and row["account"] != t.account.upper():
        return False
    rates = [x for x in (row["market"], row["net"]) if x is not None]
    return any(abs(x - t.rate) <= 0.011 for x in rates)


class _Clock:
    def __init__(self):
        self.t0 = self.t = time.time()
        self.steps: list[list] = []

    def lap(self, name: str) -> None:
        now = time.time()
        self.steps.append([name, round(now - self.t, 1)])
        self.t = now

    def total(self) -> float:
        return round(time.time() - self.t0, 1)


def _phase_upload(job_id: int, b: BrowserSession, data: bytes, file_name: str, rows: list[TxnRow]) -> None:
    s = get_settings()
    clock = _Clock()
    accounts = sorted({r.account for r in rows})
    result: dict = {"rows": [asdict(r) for r in rows], "accounts": {}}

    _save(job_id, step="Checking Equity Staging before the upload")
    b._goto(s.base_url + s.finesse_staging_route, "Open Equity Staging")
    before = {a: _staging_search(b, a) for a in accounts}
    clock.lap("Staging checked before upload")

    _save(job_id, step="Uploading the file (Equity Uploads > Transactions Upload)")
    result["upload_message"] = _upload_file(b, data, file_name)
    clock.lap("File uploaded")

    _save(job_id, step="Finding the uploaded rows in Equity Staging")
    b._goto(s.base_url + s.finesse_staging_route, "Open Equity Staging")
    for a in accounts:
        mine = [r for r in rows if r.account == a]
        all_rows, mapped = _views(b, _staging_search(b, a))
        cands = []
        for k, t in enumerate(mine):
            n_before = sum(1 for x in before.get(a, []) if _matches(x, t))
            n_all = sum(1 for x in all_rows if _matches(x, t))
            n_mapped = sum(1 for x in mapped if _matches(x, t))
            if n_before:
                kind = "duplicate"     # an identical row was already waiting: not ticked by default
            elif n_all <= 0:
                kind = "not_found"
            elif n_mapped <= 0:
                kind = "unmapped"
            else:
                kind = "new"
            cands.append({"id": f"{a}#{k}", "kind": kind, "row": asdict(t), "label": t.label()})
        result["accounts"][a] = {"client": mine[0].client, "candidates": cands}
    clock.lap("Uploaded rows found in staging")
    result["timings"] = {"upload_phase": clock.steps, "upload_total": clock.total()}

    kinds = [c["kind"] for v in result["accounts"].values() for c in v["candidates"]]
    new, dup = kinds.count("new"), kinds.count("duplicate")
    other = len(kinds) - new - dup
    msg = (f"{new} new row(s) ready to post" + (f"; {dup} already in staging (unticked)" if dup else "")
           + (f"; {other} cannot be posted (unmapped / not found)" if other else "")
           + f" — found in {clock.total():.0f}s")
    _save(job_id, status="awaiting_confirmation" if (new or dup) else "failed", step="Waiting for your confirmation",
          message=msg if (new or dup) else "Nothing to post: " + msg, data=result)


def _phase_post(job_id: int, b: BrowserSession, data: dict, selected: list[str]) -> None:
    s = get_settings()
    clock = _Clock()
    outcome: dict = {}
    chosen = set(selected)
    if "staging" not in b.page.url:
        b._goto(s.base_url + s.finesse_staging_route, "Open Equity Staging")
    for a, info in data["accounts"].items():
        plan = [TxnRow(**c["row"]) for c in info["candidates"] if c["id"] in chosen]
        if not plan:
            continue
        _save(job_id, step=f"Posting {len(plan)} row(s) for {a}")
        try:
            _, rows = _views(b, _staging_search(b, a))
            picked: list[int] = []
            for t in plan:
                pos = next((x["pos"] for x in rows if x["pos"] not in picked and _matches(x, t)), None)
                if pos is None:
                    raise PostError(f"row {t.label()} is no longer in Equity Staging — nothing posted for {a}.")
                picked.append(pos)
            for pos in picked:
                cb = b.page.locator(f'[data-cne-pos="{pos}"] input[type=checkbox]').first
                if cb.count() == 0:
                    raise b._report_error(f"{a}: no tick box on the staging row.")
                if not cb.is_checked():
                    cb.check(force=True)
            ticked = b.page.evaluate("""() => [...document.querySelectorAll('table tbody input[type=checkbox]')]
                                                .filter(c => c.checked).length""")
            if ticked != len(picked):
                raise b._report_error(f"{a}: expected {len(picked)} ticked row(s) but {ticked} are ticked — not posting.")
            if not _click_button(b, r"^\s*post\s*$", exclude=r"^$", learn_key="post_button"):
                raise b._report_error("'Post' button not found on Equity Staging.")
            _close_dialog(b, ("Yes", "Confirm", "Yes, Post", "Post", "OK", "Ok"), wait_s=5,
                          learn_key="post_confirm")                     # "Are you sure?"
            msg = _wait_message(b, 60)
            if re.search(r"error|fail", msg, re.I) and not re.search(r"success|posted", msg, re.I):
                raise b._report_error(f"{a}: Finesse reported a problem while posting: {msg}")
            _, left = _views(b, _staging_search(b, a))
            sigs: dict[tuple, list] = {}
            for t in plan:
                sigs.setdefault(t.key(), [t, 0])[1] += 1
            posted = sum(min(n, sum(_matches(x, t) for x in rows) - sum(_matches(x, t) for x in left))
                         for t, n in sigs.values())
            posted = max(posted, 0)
            outcome[a] = {"posted": posted, "selected": len(plan), "message": msg,
                          "still_in_staging": len(plan) - posted}
        except (PostError, LayoutChangedError) as e:
            outcome[a] = {"posted": 0, "selected": len(plan), "message": redact(str(e))[:300],
                          "still_in_staging": len(plan)}
        clock.lap(f"Posted {a}")
    posted = sum(v["posted"] for v in outcome.values())
    planned = sum(v["selected"] for v in outcome.values())
    data["post_result"] = outcome
    data["selected"] = selected
    data.setdefault("timings", {})["post_phase"] = clock.steps
    data["timings"]["post_total"] = clock.total()
    ok = posted == planned
    _save(job_id, status="posted" if ok else "partial", step="Done", data=data,
          message=f"Posted {posted} of {planned} selected row(s) in Finesse in {clock.total():.0f}s"
          + ("" if ok else " — check Equity Staging"))
